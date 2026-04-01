"""
本地 Excel 文件服务
支持 xlsx 文件的上传、读取、导出、下载和删除操作
"""
import os
import io
import uuid
from typing import Dict, Any, List, Optional
from pathlib import Path
from datetime import datetime

try:
    import openpyxl
    from openpyxl import Workbook, load_workbook
    from openpyxl.utils import get_column_letter
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

from core.logger import logger
from core.exceptions import ValidationError


class LocalExcelService:
    """
    本地 Excel 文件服务
    
    提供本地 xlsx 文件的上传、解析、读取、导出、下载和删除功能
    """
    
    UPLOAD_DIR = Path(__file__).parent.parent / "data" / "uploads"
    EXPORT_DIR = Path(__file__).parent.parent / "data" / "exports"
    
    def __init__(self):
        """初始化本地 Excel 服务，确保数据目录存在"""
        self.UPLOAD_DIR.mkdir(parents=True, exist_ok=True)
        self.EXPORT_DIR.mkdir(parents=True, exist_ok=True)
        self._file_cache: Dict[str, Dict[str, Any]] = {}
    
    def _check_openpyxl(self) -> None:
        """
        检查 openpyxl 是否可用
        
        Raises:
            ValidationError: openpyxl 未安装
        """
        if not OPENPYXL_AVAILABLE:
            raise ValidationError("openpyxl 未安装，无法处理 xlsx 文件")
    
    def upload_file(self, file_content: bytes, filename: str) -> Dict[str, Any]:
        """
        上传并解析 xlsx 文件
        
        Args:
            file_content: 文件内容字节流
            filename: 原始文件名
            
        Returns:
            文件信息字典，包含 file_id、filename、sheets、row_count、column_count
            
        Raises:
            ValidationError: 文件格式不支持或解析失败
        """
        self._check_openpyxl()
        
        if not filename.lower().endswith('.xlsx'):
            raise ValidationError("仅支持 .xlsx 格式文件")
        
        file_id = f"local_{uuid.uuid4().hex[:12]}"
        file_path = self.UPLOAD_DIR / f"{file_id}.xlsx"
        
        with open(file_path, 'wb') as f:
            f.write(file_content)
        
        try:
            wb = load_workbook(file_path, read_only=True)
            sheets = wb.sheetnames
            
            ws = wb.active
            row_count = ws.max_row or 0
            col_count = ws.max_column or 0
            
            wb.close()
            
            self._file_cache[file_id] = {
                "path": str(file_path),
                "filename": filename,
                "uploaded_at": datetime.now()
            }
            
            logger.info(f"上传文件成功: {file_id}, {filename}")
            
            return {
                "file_id": file_id,
                "filename": filename,
                "sheets": sheets,
                "row_count": row_count,
                "column_count": col_count
            }
            
        except Exception as e:
            file_path.unlink(missing_ok=True)
            logger.error(f"解析文件失败: {e}")
            raise ValidationError(f"无法解析 xlsx 文件: {str(e)}")
    
    def read_file(
        self,
        file_id: str,
        sheet_name: Optional[str] = None,
        range_str: Optional[str] = None
    ) -> Dict[str, Any]:
        """
        读取已上传的 xlsx 文件
        
        Args:
            file_id: 文件 ID
            sheet_name: 工作表名称（可选）
            range_str: 读取范围，如 "A1:C10"（可选）
            
        Returns:
            表格数据字典，包含 file_id、sheet_name、range、values、row_count、column_count
            
        Raises:
            ValidationError: 文件不存在或工作表不存在
        """
        self._check_openpyxl()
        
        file_info = self._file_cache.get(file_id)
        if not file_info:
            file_path = self.UPLOAD_DIR / f"{file_id}.xlsx"
            if not file_path.exists():
                raise ValidationError(f"文件不存在: {file_id}")
            file_info = {"path": str(file_path)}
        
        wb = load_workbook(file_info["path"], read_only=True, data_only=True)
        
        try:
            if sheet_name:
                if sheet_name not in wb.sheetnames:
                    raise ValidationError(f"工作表不存在: {sheet_name}")
                ws = wb[sheet_name]
            else:
                ws = wb.active
                sheet_name = ws.title
            
            if range_str:
                values = self._read_range(ws, range_str)
            else:
                values = self._read_all(ws)
            
            return {
                "file_id": file_id,
                "sheet_name": sheet_name,
                "range": range_str or f"A1:{get_column_letter(ws.max_column or 1)}{ws.max_row or 1}",
                "values": values,
                "row_count": len(values),
                "column_count": len(values[0]) if values else 0
            }
            
        finally:
            wb.close()
    
    def _read_range(self, ws, range_str: str) -> List[List[Any]]:
        """
        读取指定范围的数据
        
        Args:
            ws: openpyxl 工作表对象
            range_str: 范围字符串，如 "A1:C10"
            
        Returns:
            二维数组形式的单元格值
        """
        values = []
        for row in ws[range_str]:
            row_values = []
            for cell in row:
                row_values.append(cell.value)
            values.append(row_values)
        return values
    
    def _read_all(self, ws) -> List[List[Any]]:
        """
        读取工作表所有数据
        
        Args:
            ws: openpyxl 工作表对象
            
        Returns:
            二维数组形式的单元格值
        """
        values = []
        for row in ws.iter_rows():
            row_values = []
            for cell in row:
                row_values.append(cell.value)
            values.append(row_values)
        return values
    
    def export_to_file(
        self,
        data: List[List[Any]],
        filename: Optional[str] = None,
        sheet_name: str = "Sheet1"
    ) -> Dict[str, Any]:
        """
        将数据导出为 xlsx 文件
        
        Args:
            data: 二维数组数据
            filename: 导出文件名（可选）
            sheet_name: 工作表名称，默认为 "Sheet1"
            
        Returns:
            导出文件信息，包含 file_id、filename、path、row_count、column_count
        """
        self._check_openpyxl()
        
        if not filename:
            filename = f"export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        elif not filename.lower().endswith('.xlsx'):
            filename += '.xlsx'
        
        wb = Workbook()
        ws = wb.active
        ws.title = sheet_name
        
        for row_idx, row_data in enumerate(data, 1):
            for col_idx, value in enumerate(row_data, 1):
                ws.cell(row=row_idx, column=col_idx, value=value)
        
        file_id = f"export_{uuid.uuid4().hex[:12]}"
        file_path = self.EXPORT_DIR / f"{file_id}.xlsx"
        wb.save(file_path)
        wb.close()
        
        logger.info(f"导出文件成功: {file_id}, {filename}")
        
        return {
            "file_id": file_id,
            "filename": filename,
            "path": str(file_path),
            "row_count": len(data),
            "column_count": len(data[0]) if data else 0
        }
    
    def get_export_file_path(self, file_id: str) -> Optional[Path]:
        """
        获取导出文件路径
        
        Args:
            file_id: 文件 ID
            
        Returns:
            文件路径，如果文件不存在则返回 None
        """
        file_path = self.EXPORT_DIR / f"{file_id}.xlsx"
        if file_path.exists():
            return file_path
        return None
    
    def delete_file(self, file_id: str) -> bool:
        """
        删除文件
        
        Args:
            file_id: 文件 ID
            
        Returns:
            删除成功返回 True，文件不存在返回 False
        """
        upload_path = self.UPLOAD_DIR / f"{file_id}.xlsx"
        if upload_path.exists():
            upload_path.unlink()
            self._file_cache.pop(file_id, None)
            logger.info(f"删除上传文件: {file_id}")
            return True
        
        export_path = self.EXPORT_DIR / f"{file_id}.xlsx"
        if export_path.exists():
            export_path.unlink()
            logger.info(f"删除导出文件: {file_id}")
            return True
        
        return False
    
    def list_files(self) -> Dict[str, List[Dict[str, Any]]]:
        """
        列出所有文件
        
        Returns:
            包含 uploads 和 exports 两个列表的字典，每个元素包含文件信息
        """
        uploads = []
        for f in self.UPLOAD_DIR.glob("*.xlsx"):
            uploads.append({
                "file_id": f.stem,
                "filename": f.name,
                "size": f.stat().st_size,
                "created_at": datetime.fromtimestamp(f.stat().st_ctime).isoformat()
            })
        
        exports = []
        for f in self.EXPORT_DIR.glob("*.xlsx"):
            exports.append({
                "file_id": f.stem,
                "filename": f.name,
                "size": f.stat().st_size,
                "created_at": datetime.fromtimestamp(f.stat().st_ctime).isoformat()
            })
        
        return {
            "uploads": uploads,
            "exports": exports
        }
